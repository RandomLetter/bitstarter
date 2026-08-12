#!/usr/bin/env python3
"""Build the Santiago 'How to Play' slide deck (PPTX).

Prerequisites (see README.md): assets/icons via render_emoji.js + make_icons.py,
assets/diagrams via make_diagrams.py. Output: Santiago-How-to-Play.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ICO = os.path.join(HERE, 'assets', 'icons')
DIA = os.path.join(HERE, 'assets', 'diagrams')

SAND = RGBColor(0xF4, 0xE8, 0xC8)
SAND_DEEP = RGBColor(0xEA, 0xD6, 0xA4)
INK = RGBColor(0x3A, 0x2C, 0x1A)
ACCENT = RGBColor(0xB0, 0x62, 0x2D)
WATER = RGBColor(0x2A, 0x7F, 0xB8)
GREEN = RGBColor(0x2E, 0x7D, 0x4F)
RED = RGBColor(0xB0, 0x3A, 0x2E)
PURPLE = RGBColor(0x80, 0x50, 0xA8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD = RGBColor(0xFF, 0xFA, 0xF0)

W, H = Inches(13.333), Inches(7.5)
CROPS = [('banana', 'Bananas'), ('sugar', 'Sugar cane'), ('potato', 'Potatoes'),
         ('bean', 'Beans'), ('pepper', 'Red peppers')]

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]


def bg(slide, color=SAND):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.text_frame.word_wrap = True
    return tb


def set_run(run, text, size, color=INK, bold=False, italic=False, font='Georgia'):
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font


def title_bar(slide, text, kicker=None):
    bar = slide.shapes.add_shape(1, 0, 0, W, Inches(1.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    bar.shadow.inherit = False
    tb = textbox(slide, Inches(0.55), Inches(0.1), Inches(12.2), Inches(0.95))
    p = tb.text_frame.paragraphs[0]
    set_run(p.add_run(), text, 29, WHITE, bold=True)
    if kicker:
        p2 = tb.text_frame.add_paragraph()
        set_run(p2.add_run(), kicker, 13, RGBColor(0xFF, 0xE8, 0xC8), italic=True)


def bullets(slide, items, x=Inches(0.55), y=Inches(1.45), w=Inches(7.2), h=Inches(5.6),
            size=17, gap=10):
    tb = textbox(slide, x, y, w, h)
    tf = tb.text_frame
    first = True
    for item in items:
        text, opts = item if isinstance(item, tuple) else (item, {})
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(opts.get('gap', gap))
        p.level = opts.get('level', 0)
        marker = '' if opts.get('nobullet') else ('– ' if p.level else '• ')
        if marker:
            set_run(p.add_run(), marker, opts.get('size', size),
                    ACCENT if p.level == 0 else INK, bold=True)
        if opts.get('lead'):
            set_run(p.add_run(), opts['lead'] + ' ', opts.get('size', size), INK, bold=True)
        set_run(p.add_run(), text, opts.get('size', size), opts.get('color', INK),
                bold=opts.get('bold', False), italic=opts.get('italic', False))
    return tb


def picture(slide, path, x, y, max_w, max_h, caption=None, border=False):
    img = Image.open(path)
    ratio = img.width / img.height
    w = max_w
    h = Emu(int(w / ratio))
    if h > max_h:
        h = max_h
        w = Emu(int(h * ratio))
    px = x + Emu(int((max_w - w) / 2))
    pic = slide.shapes.add_picture(path, px, y, w, h)
    if border:
        pic.line.color.rgb = INK
        pic.line.width = Pt(1.5)
    if caption:
        tb = textbox(slide, x, y + h + Emu(50000), max_w, Inches(0.7))
        p = tb.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        set_run(p.add_run(), caption, 12.5, INK, italic=True)
    return pic


def diagram(slide, name, x, y, max_w, max_h, caption=None):
    return picture(slide, f'{DIA}/{name}.png', x, y, max_w, max_h, caption=caption)


def icon(slide, name, x, y, size):
    return slide.shapes.add_picture(f'{ICO}/{name}.png', x, y, size, size)


def label_under(slide, text, x, y, w, size=13, bold=False, color=INK):
    tb = textbox(slide, x, y, w, Inches(0.45))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), text, size, color, bold=bold)
    return tb


def crop_row(slide, x, y, tile=Inches(1.15), gap=Inches(0.45), label_size=13):
    cx = x
    for key, name in CROPS:
        icon(slide, f'tile_{key}', cx, y, tile)
        if label_size:
            label_under(slide, name, cx - Inches(0.2), y + tile + Inches(0.04),
                        tile + Inches(0.4), size=label_size)
        cx += tile + gap


def card(slide, x, y, w, h):
    c = slide.shapes.add_shape(1, x, y, w, h)
    c.fill.solid()
    c.fill.fore_color.rgb = CARD
    c.line.color.rgb = RGBColor(0xD8, 0xC0, 0x8A)
    c.line.width = Pt(1.25)
    c.shadow.inherit = False
    return c


def big_text(slide, text, x, y, w, size, color=INK, bold=True, align=PP_ALIGN.CENTER):
    tb = textbox(slide, x, y, w, Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size, color, bold=bold)
    return tb


def footer(slide, n):
    tb = textbox(slide, Inches(11.9), Inches(7.05), Inches(1.2), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    set_run(p.add_run(), f'Santiago · {n}', 10, RGBColor(0x8A, 0x70, 0x4A))


def new_slide():
    s = prs.slides.add_slide(BLANK)
    bg(s)
    return s


# ---- 1. Title -----------------------------------------------------------------
s = new_slide()
bg(s, SAND_DEEP)
tb = textbox(s, Inches(0.7), Inches(0.95), Inches(11.9), Inches(3.4))
p = tb.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
set_run(p.add_run(), 'SANTIAGO', 66, ACCENT, bold=True)
p = tb.text_frame.add_paragraph()
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(6)
set_run(p.add_run(), '“The water flows where the money goes!”', 25, INK, italic=True, bold=True)
p = tb.text_frame.add_paragraph()
p.alignment = PP_ALIGN.CENTER
p.space_before = Pt(20)
set_run(p.add_run(),
        'A game of plantations, irrigation and bribery — Claudia Hely & Roman Pelek (AMIGO, 2003)',
        16, INK)
p = tb.text_frame.add_paragraph()
p.alignment = PP_ALIGN.CENTER
set_run(p.add_run(), '3–5 players · ages 10+ · about 60 minutes', 16, INK)
crop_row(s, Inches(2.75), Inches(4.75))

# ---- 2. The story & the goal -----------------------------------------------------
s = new_slide()
title_bar(s, 'The island of Santiago', 'What the game is about')
bullets(s, [
    ('On the hot Cape Verde island of Santiago, every drop of water is precious.', {}),
    ('Buy plantations at auction, join them into big same-crop areas, and get canals built to them — dry plantations crumble to desert.', {'lead': 'Your goal:'}),
    ('Canals go where the Canal Overseer wants. The Overseer listens to bribes.', {'lead': 'The twist:'}),
    ('Most money at the end wins. Every Escudo spent must earn itself back.', {'lead': 'Winning:'}),
], w=Inches(12.3), size=19, gap=16)
big_text(s, 'Five crops to grow:', Inches(0.55), Inches(4.55), Inches(12.2), 16, ACCENT)
crop_row(s, Inches(3.05), Inches(5.15))
footer(s, 2)

# ---- 3. Components -----------------------------------------------------------------
s = new_slide()
title_bar(s, 'What’s in the box')
card(s, Inches(0.55), Inches(1.5), Inches(6.0), Inches(2.6))
big_text(s, '45 plantation tiles', Inches(0.75), Inches(1.65), Inches(5.6), 18, ACCENT)
crop_row(s, Inches(0.95), Inches(2.25), tile=Inches(0.92), gap=Inches(0.18), label_size=11)
label_under(s, '9 per crop — with 1 or 2 planters in the corner', Inches(0.75), Inches(3.62), Inches(5.6), 12)
card(s, Inches(6.85), Inches(1.5), Inches(5.9), Inches(2.6))
big_text(s, '22 yield markers each', Inches(7.05), Inches(1.65), Inches(5.5), 18, ACCENT)
s.shapes.add_picture(f'{ICO}/markers_row.png', Inches(7.95), Inches(2.45), Inches(3.7), Inches(0.66))
label_under(s, 'your stake in a plantation — one colour per player', Inches(7.05), Inches(3.45), Inches(5.5), 12)
card(s, Inches(0.55), Inches(4.35), Inches(6.0), Inches(2.6))
big_text(s, '15 canals · 1 spring', Inches(0.75), Inches(4.5), Inches(5.6), 18, ACCENT)
diagram(s, 'canal_pieces', Inches(1.3), Inches(5.25), Inches(4.5), Inches(1.3))
card(s, Inches(6.85), Inches(4.35), Inches(5.9), Inches(2.6))
big_text(s, 'Escudos & the Canal Overseer', Inches(7.05), Inches(4.5), Inches(5.5), 18, ACCENT)
icon(s, 'money', Inches(8.15), Inches(5.15), Inches(1.1))
icon(s, 'overseer', Inches(10.0), Inches(5.15), Inches(1.1))
label_under(s, 'start with 10 — keep yours hidden!', Inches(7.0), Inches(6.3), Inches(2.9), 12)
label_under(s, 'decides where water flows', Inches(9.4), Inches(6.3), Inches(2.6), 12)
footer(s, 3)

# ---- 4. Setup -------------------------------------------------------------------------
s = new_slide()
title_bar(s, 'Setting up')
bullets(s, [
    ('Spring on any ditch intersection — the middle is friendlier.', {'lead': 'Spring:'}),
    ('3–4 players: four face-down piles of 11 tiles (one tile removed). 5 players: five piles of 9.', {'lead': 'Tiles:'}),
    ('11 canals by the board (9 with five players) + one personal “extra canal” each.', {'lead': 'Canals:'}),
    ('10 Escudos, 22 markers and a proposal canal per player; someone random takes the Overseer figure.', {'lead': 'Players:'}),
    ('11 rounds (3–4 players) or 9 rounds (5) — exactly until all tiles are placed.', {'lead': 'Length:'}),
], w=Inches(6.3), size=17, gap=14)
diagram(s, 'setup', Inches(7.1), Inches(1.9), Inches(5.7), Inches(4.2),
        caption='Board with the spring on a central intersection; face-down tile piles beside it')
footer(s, 4)

# ---- 5. Canals -------------------------------------------------------------------------
s = new_slide()
title_bar(s, 'Key concept — building canals')
bullets(s, [
    ('Canals lie on the brown ditch lines, always intersection to intersection.', {}),
    ('The network grows out from the spring — every new canal must touch the spring or an existing canal.', {}),
    ('Branching and loops are fine.', {}),
], w=Inches(6.4), size=18, gap=14)
icon(s, 'water', Inches(0.85), Inches(4.5), Inches(0.9))
big_text(s, 'Only 15 canals all game —\nwater is the scarcest resource', Inches(1.95), Inches(4.55),
         Inches(4.9), 17, WATER, align=PP_ALIGN.LEFT)
diagram(s, 'canals', Inches(7.2), Inches(1.55), Inches(5.6), Inches(4.7),
        caption='✔ touches the network — allowed.   ✘ floats free of the network — not allowed.')
footer(s, 5)

# ---- 6. Irrigation & yield markers ------------------------------------------------------
s = new_slide()
title_bar(s, 'Key concept — irrigation & yield markers')
bullets(s, [
    ('Place a plantation, add your markers: one per planter shown on the tile.', {}),
    ('A plantation touching a canal on any side is irrigated — forever.', {'lead': 'Irrigated:'}),
], w=Inches(12.3), size=18, gap=10)
card(s, Inches(0.7), Inches(3.0), Inches(5.6), Inches(3.7))
icon(s, 'diag_bid1', Inches(1.25), Inches(3.35), Inches(1.9))
big_text(s, '1 planter → 1 marker', Inches(0.85), Inches(5.45), Inches(2.7), 14)
icon(s, 'diag_bid2', Inches(3.85), Inches(3.35), Inches(1.9))
big_text(s, '2 planters → 2 markers', Inches(3.45), Inches(5.45), Inches(2.7), 14)
label_under(s, 'markers = your share of the harvest', Inches(0.85), Inches(6.15), Inches(5.3), 13, bold=True)
diagram(s, 'irrigation', Inches(6.9), Inches(3.0), Inches(6.0), Inches(3.3),
        caption='Next to a canal: irrigated ✔ — cut off from the water: drying ✘')
footer(s, 6)

# ---- 7. Connected areas --------------------------------------------------------------------
s = new_slide()
title_bar(s, 'Key concept — connected areas', 'Where the money comes from')
bullets(s, [
    ('Same-crop tiles touching side by side form one area — markers of several players can share it.', {}),
    ('Diagonal never counts; a canal running through does NOT split an area.', {}),
    ('Area value for you = its tiles × your markers on it.', {'lead': 'Value:'}),
], w=Inches(6.4), size=18, gap=12)
big_text(s, 'The 4-tile banana area is worth:', Inches(0.55), Inches(4.35), Inches(6.2), 15, INK)
big_text(s, 'Beige: 4 tiles × 3 markers = 12 Esc', Inches(0.55), Inches(4.82), Inches(6.2), 17, ACCENT)
big_text(s, 'Purple: 4 tiles × 1 marker = 4 Esc', Inches(0.55), Inches(5.3), Inches(6.2), 17, PURPLE)
big_text(s, 'The diagonal banana belongs to a different area.', Inches(0.55), Inches(5.9),
         Inches(6.2), 14, RED, bold=False)
diagram(s, 'areas', Inches(7.2), Inches(1.55), Inches(5.6), Inches(4.7),
        caption='One 4-tile banana area (canal through it and a neutral tile included) — the ✘ tile touches only diagonally')
footer(s, 7)

# ---- 8. Round overview -----------------------------------------------------------------------
s = new_slide()
title_bar(s, 'A round at a glance', 'Seven phases, always in this order')
phases = [
    ('1', 'cards', 'Reveal & bid', 'one open bid each — or pass'),
    ('2', 'overseer', 'New Overseer', 'lowest bid (or first pass) takes the figure'),
    ('3', 'seedling', 'Take & place tiles', 'highest bidder picks first'),
    ('4', 'handshake', 'Bribe the Overseer', 'one canal gets built'),
    ('5', 'water', 'Extra irrigation', 'one player may use their extra canal'),
    ('6', 'sun', 'Drying', 'dry plantations suffer (skip in last round)'),
    ('7', 'coin', 'Income', 'everyone collects 3 Escudos (skip in last round)'),
]
y = Inches(1.4)
for num, ic, name, desc in phases:
    chip = s.shapes.add_shape(9, Inches(0.6), y, Inches(0.48), Inches(0.48))
    chip.fill.solid()
    chip.fill.fore_color.rgb = WATER if num in '45' else ACCENT
    chip.line.fill.background()
    chip.shadow.inherit = False
    cp = chip.text_frame.paragraphs[0]
    cp.alignment = PP_ALIGN.CENTER
    set_run(cp.add_run(), num, 17, WHITE, bold=True)
    icon(s, ic, Inches(1.3), y - Inches(0.02), Inches(0.52))
    tb = textbox(s, Inches(2.05), y - Inches(0.03), Inches(10.7), Inches(0.7))
    p = tb.text_frame.paragraphs[0]
    set_run(p.add_run(), name + '  ', 18, INK, bold=True)
    set_run(p.add_run(), '— ' + desc, 16, INK)
    y += Inches(0.79)
footer(s, 8)

# ---- 9. Auction --------------------------------------------------------------------------------
s = new_slide()
title_bar(s, 'Phases 1–2 — the auction', 'One bid each, no second chances')
bullets(s, [
    ('Flip one tile per pile, then everyone bids once, openly — or passes.', {}),
    ('Every bid must be a NEW amount — bidding lower is allowed. All bids go to the bank, win or lose.', {'lead': 'Unique bids:'}),
    ('Lowest bid — or the first passer — becomes the new Canal Overseer.', {'lead': 'Consolation:'}),
], w=Inches(12.3), size=18, gap=12)
# worked example as cards: bid → outcome
example = [
    ('Bernd', '5 Esc', 'picks first', GREEN, None),
    ('Dagmar', '4 Esc', 'picks 2nd', INK, None),
    ('Anika', '1 Esc', 'picks 3rd', INK, None),
    ('Chris', 'PASS', 'picks last', RED, 'overseer'),
]
x = Inches(0.7)
for name, bid, outcome, color, extra in example:
    card(s, x, Inches(3.85), Inches(2.85), Inches(2.5))
    big_text(s, name, x, Inches(4.0), Inches(2.85), 17)
    big_text(s, bid, x, Inches(4.55), Inches(2.85), 24, color)
    big_text(s, outcome, x, Inches(5.25), Inches(2.85), 14, INK, bold=False)
    if extra:
        icon(s, extra, x + Inches(1.13), Inches(5.55), Inches(0.6))
        big_text(s, 'new Overseer!', x, Inches(6.12), Inches(2.85), 13, ACCENT)
    x += Inches(3.05)
big_text(s, 'High bid = first pick of land.  Low bid = power over the water.',
         Inches(0.55), Inches(6.75), Inches(12.2), 15, WATER)
footer(s, 9)

# ---- 10. Placement -------------------------------------------------------------------------------
s = new_slide()
title_bar(s, 'Phase 3 — take & place plantations')
bullets(s, [
    ('In descending bid order: take one revealed tile, place it on any free square, add your markers.', {}),
    ('Passed players place one marker fewer than the planters shown.', {'lead': 'Penalty:'}),
    ('With 3 players the leftover tile is placed neutral — no markers — next to another plantation.', {'lead': '3 players:'}),
], w=Inches(12.3), size=18, gap=12)
card(s, Inches(1.5), Inches(3.6), Inches(4.8), Inches(3.2))
big_text(s, 'You bid', Inches(1.5), Inches(3.75), Inches(4.8), 17, GREEN)
icon(s, 'diag_bid2', Inches(2.95), Inches(4.3), Inches(1.9))
big_text(s, '2 planters → your 2 markers', Inches(1.5), Inches(6.3), Inches(4.8), 14)
card(s, Inches(7.0), Inches(3.6), Inches(4.8), Inches(3.2))
big_text(s, 'You passed', Inches(7.0), Inches(3.75), Inches(4.8), 17, RED)
icon(s, 'diag_pass2', Inches(8.45), Inches(4.3), Inches(1.9))
big_text(s, '2 planters → only 1 marker', Inches(7.0), Inches(6.3), Inches(4.8), 14)
footer(s, 10)

# ---- 11. Bribery -----------------------------------------------------------------------------------
s = new_slide()
title_bar(s, 'Phase 4 — bribe the Canal Overseer', 'The heart of the game')
opts = [
    ('handshake', 'Propose', 'a canal + a bribe of 1+ Esc, placed openly'),
    ('coin', 'Support', 'someone’s proposal with 1+ Esc more'),
    ('timer', 'Pass', 'and keep your money'),
]
y = Inches(1.5)
for ic, name, desc in opts:
    card(s, Inches(0.55), y, Inches(6.2), Inches(1.0))
    icon(s, ic, Inches(0.75), y + Inches(0.17), Inches(0.65))
    tb = textbox(s, Inches(1.6), y + Inches(0.12), Inches(5.0), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    set_run(p.add_run(), name + ' ', 17, ACCENT, bold=True)
    set_run(p.add_run(), desc, 15, INK)
    y += Inches(1.15)
dec = card(s, Inches(0.55), Inches(5.0), Inches(6.2), Inches(1.9))
dec.fill.fore_color.rgb = RGBColor(0xE2, 0xF0, 0xFA)
tb = textbox(s, Inches(0.75), Inches(5.1), Inches(5.9), Inches(1.75))
p = tb.text_frame.paragraphs[0]
set_run(p.add_run(), 'The Overseer then either ', 15, INK)
set_run(p.add_run(), 'accepts ONE proposal', 15, WATER, bold=True)
set_run(p.add_run(), ' — any one! — and pockets its bribe…', 15, INK)
p2 = tb.text_frame.add_paragraph()
p2.space_before = Pt(8)
set_run(p2.add_run(), '…or ', 15, INK)
set_run(p2.add_run(), 'builds anywhere else', 15, WATER, bold=True)
set_run(p2.add_run(), ', paying the bank 1 Esc more than the highest bribe. Losing bribes are returned.', 15, INK)
diagram(s, 'bribe', Inches(7.2), Inches(1.7), Inches(5.6), Inches(4.4),
        caption='Two proposals (dashed, with bribes): Purple wants the beans watered, Gray the potatoes')
footer(s, 11)

# ---- 12. Phases 5-7 -----------------------------------------------------------------------------------
s = new_slide()
title_bar(s, 'Phases 5–7 — extra canal, drying, income')
rows = [
    ('water', 'Extra irrigation', 'your ONE free canal for the whole game — only one player may build it per round'),
    ('sun', 'Drying', 'every plantation without a canal loses a marker; none left → desert'),
    ('coin', 'Income', 'everyone collects 3 Escudos'),
]
y = Inches(1.45)
for ic, name, desc in rows:
    icon(s, ic, Inches(0.7), y, Inches(0.7))
    tb = textbox(s, Inches(1.65), y - Inches(0.02), Inches(10.9), Inches(1.1))
    p = tb.text_frame.paragraphs[0]
    set_run(p.add_run(), name + '  ', 18, ACCENT, bold=True)
    set_run(p.add_run(), '— ' + desc, 15, INK)
    y += Inches(0.95)
big_text(s, 'A plantation left dry, round after round:', Inches(0.55), Inches(4.45), Inches(12.2), 16, INK)
diagram(s, 'drying_seq', Inches(2.4), Inches(5.05), Inches(8.5), Inches(1.6),
        caption='…and desert can never be cultivated again. (Drying and income are skipped in the last round.)')
footer(s, 12)

# ---- 13. Scoring -----------------------------------------------------------------------------------------
s = new_slide()
title_bar(s, 'Game end & scoring', 'After the last tile is placed')
bullets(s, [
    ('Everything still dry turns to desert — even with markers on it.', {'lead': 'Final drought:'}),
    ('Keep your cash, then every area pays every player:', {'lead': 'Score:'}),
], w=Inches(6.4), size=18, gap=12)
fx = Inches(0.55)
fy = Inches(3.15)
for txt, color in [('tiles in area', GREEN), ('×', INK), ('your markers', ACCENT), ('=', INK), ('Escudos', WATER)]:
    wbox = Inches(1.75) if txt not in '×=' else Inches(0.4)
    if txt in '×=':
        big_text(s, txt, fx, fy + Inches(0.18), wbox, 24)
    else:
        card(s, fx, fy, wbox, Inches(0.95))
        big_text(s, txt, fx, fy + Inches(0.28), wbox, 16, color)
    fx += wbox + Inches(0.1)
big_text(s, 'Beige: 4 tiles × 4 markers = 16 Esc', Inches(0.55), Inches(4.5),
         Inches(6.4), 16, ACCENT, align=PP_ALIGN.LEFT)
big_text(s, 'Gray: 4 tiles × 1 marker = 4 Esc', Inches(0.55), Inches(4.98),
         Inches(6.4), 16, INK, align=PP_ALIGN.LEFT)
big_text(s, 'The desert tile pays nobody.', Inches(0.55), Inches(5.46),
         Inches(6.4), 14, RED, bold=False, align=PP_ALIGN.LEFT)
icon(s, 'trophy', Inches(0.85), Inches(6.0), Inches(0.9))
big_text(s, 'Most money wins!', Inches(1.95), Inches(6.18), Inches(4.5), 22, ACCENT,
         align=PP_ALIGN.LEFT)
diagram(s, 'scoring', Inches(7.2), Inches(1.7), Inches(5.6), Inches(4.4),
        caption='Only irrigated, connected tiles pay — the flipped desert tile is gone for good')
footer(s, 13)

# ---- 14. Tips ----------------------------------------------------------------------------------------------
s = new_slide()
title_bar(s, 'Tips for your first game')
tips = [
    ('water', 'Water is scarce', 'only 15 canals all game — many fields WILL become desert'),
    ('overseer', 'Cheap bids buy power', 'the Overseer earns bribes and steers the water'),
    ('think', 'Check who profits', 'growing an area helps everyone with markers on it'),
    ('timer', 'Save the extra canal', 'one per game — spend it where no bribe will help'),
]
positions = [(Inches(0.55), Inches(1.6)), (Inches(6.85), Inches(1.6)),
             (Inches(0.55), Inches(4.3)), (Inches(6.85), Inches(4.3))]
for (ic, head, body), (x, y) in zip(tips, positions):
    card(s, x, y, Inches(5.9), Inches(2.4))
    icon(s, ic, x + Inches(0.3), y + Inches(0.3), Inches(0.85))
    tb = textbox(s, x + Inches(1.4), y + Inches(0.28), Inches(4.3), Inches(1.9))
    p = tb.text_frame.paragraphs[0]
    set_run(p.add_run(), head, 19, ACCENT, bold=True)
    p2 = tb.text_frame.add_paragraph()
    p2.space_before = Pt(6)
    set_run(p2.add_run(), body, 15, INK)
footer(s, 14)

# ---- 15. Play online ------------------------------------------------------------------------------------------
s = new_slide()
bg(s, SAND_DEEP)
band = s.shapes.add_shape(1, 0, Inches(2.2), W, Inches(2.6))
band.fill.solid()
band.fill.fore_color.rgb = WATER
band.line.fill.background()
band.shadow.inherit = False
tb = textbox(s, Inches(0.8), Inches(2.5), Inches(11.7), Inches(2.0))
p = tb.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
set_run(p.add_run(), 'Play it online — asynchronously', 32, WHITE, bold=True)
p = tb.text_frame.add_paragraph()
p.alignment = PP_ALIGN.CENTER
set_run(p.add_run(), 'budoludo.com/santiago', 28, RGBColor(0xFF, 0xE8, 0xA8), bold=True)
icon(s, 'email', Inches(6.32), Inches(5.15), Inches(0.7))
tb = textbox(s, Inches(0.8), Inches(5.95), Inches(11.7), Inches(1.3))
for line in ['Personal secret links — no accounts, no passwords.',
             'An email lands whenever it is your turn.']:
    p = tb.text_frame.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    set_run(p.add_run(), line, 17, INK)
crop_row(s, Inches(4.3), Inches(0.6), tile=Inches(0.8), gap=Inches(0.25), label_size=0)

out = os.path.join(HERE, 'Santiago-How-to-Play.pptx')
prs.save(out)
print('saved', out, '· slides:', len(prs.slides._sldIdLst))
